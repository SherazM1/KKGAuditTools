"""Editable report content, original-photo handling, and provider-free UI tests."""

from dataclasses import asdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from fontTools.ttLib import TTFont
import subprocess
import sys
import tempfile
import textwrap
import unittest
import uuid
from xml.etree import ElementTree as ET
from unittest.mock import Mock, patch
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from PIL import Image
from streamlit.testing.v1 import AppTest

from shelf_audit.reporting import ReportMetadata, build_report
from shelf_audit.reporting.docx_renderer import (
    BLACK_FONT, BOLD_FONT, MEDIUM_FONT, LAYOUT_PRESETS, MIN_BODY_FONT_SIZE,
    PHOTO_MAX_WIDTH, PHOTO_MAX_HEIGHT, _set_font, estimate_report_height,
    render_docx, select_layout, photo_box_width, FRAME_WIDTH, FRAME_HEIGHT,
    FRAME_X, FRAME_Y, HEADER_TOP_MARGIN, BOOK_FONT, TRACKING_PT, LOGO_X, LOGO_Y, LOGO_WIDTH,
)
from shelf_audit.reporting.formatting import DEFAULT_LOGO_PATH
from shelf_audit.reporting.formatting import BLUE
from shelf_audit.reporting.templates import NO_OPPORTUNITIES_TEXT
from test_shelf_audit_reporting import completed_result


def photo_bytes(size=(600, 300), color="red", image_format="PNG"):
    output = BytesIO()
    Image.new("RGB", size, color).save(output, format=image_format)
    return output.getvalue()


def docx_text(data):
    doc = Document(BytesIO(data))
    paragraphs = []
    # Preserve visible body order when the first finding is inside a table.
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraphs.append(block)
        else:
            for row in block.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in (*section.header.tables, *section.footer.tables):
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
    return "\n".join(p.text for p in paragraphs)


def media_bytes(data):
    with ZipFile(BytesIO(data)) as package:
        return [package.read(name) for name in package.namelist() if name.startswith("word/media/")]


class ShelfAuditDocxTests(unittest.TestCase):

    def test_approved_typography_tracking_and_metadata_order(self):
        for count in (1, 2, 3, 4, 8):
            report = build_report(completed_result(count), metadata=ReportMetadata(client='Tide', retailer='Walmart', location='Bentonville, Arkansas'))
            doc = Document(BytesIO(render_docx(report)))
            self.assertEqual(doc.styles['Normal'].font.name, BOOK_FONT)
            self.assertEqual(select_layout(report), LAYOUT_PRESETS[0])
            def check(run, font, size, tracked=False):
                self.assertEqual(run.font.name, font)
                self.assertEqual(run.font.size.pt, size)
                self.assertEqual(run.bold, font in (BLACK_FONT, BOLD_FONT))
                for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    self.assertEqual(run._r.rPr.rFonts.get(qn('w:'+script)), font)
                spacing = run._r.rPr.find(qn('w:spacing'))
                if tracked:
                    self.assertEqual(spacing.get(qn('w:val')), str(round(TRACKING_PT * 20)))
            title, subtitle = doc.paragraphs[:2]
            self.assertEqual(title.text, report.title)
            check(title.runs[0], BLACK_FONT, 30)
            self.assertEqual(subtitle.text, report.subtitle.upper())
            check(subtitle.runs[0], BOLD_FONT, 11, True)
            self.assertEqual(str(subtitle.runs[0].font.color.rgb), BLUE.lstrip('#'))
            for paragraph, label, value in zip(doc.paragraphs[2:5], ('Client', 'Retailer', 'Location'), ('Tide', 'Walmart', 'Bentonville, Arkansas')):
                self.assertEqual(paragraph.text, f'{label}: {value}')
                check(paragraph.runs[0], BLACK_FONT, 12)
                check(paragraph.runs[1], BOOK_FONT, 12)
            for item in report.opportunities:
                heading = next(p for p in doc.paragraphs if p.text == f'{item.number:02d} {item.title}')
                check(heading.runs[0], BLACK_FONT, 12)
                for text in (item.evidence, item.recommendation):
                    check(next(p for p in doc.paragraphs if p.text == text).runs[0], BOOK_FONT, 10)
            for paragraph in doc.paragraphs:
                if paragraph.text in ('WHY', 'RECOMMENDATION'):
                    check(paragraph.runs[0], BOLD_FONT, 11, True)
                    self.assertEqual(str(paragraph.runs[0].font.color.rgb), BLUE.lstrip('#'))

    def test_reference_anchors_logo_panel_caption_and_original_assets(self):
        for photo in (photo_bytes((300,900)), photo_bytes((900,300)), None, b'invalid'):
            data = render_docx(build_report(completed_result(3), audit_photo=photo))
            doc = Document(BytesIO(data))
            logo = doc.sections[0].header._element.xpath('//wp:anchor')[0]
            def position(anchor, axis):
                return int(anchor.find(qn('wp:position'+axis)).find(qn('wp:posOffset')).text) / 914400
            self.assertAlmostEqual(position(logo, 'H'), LOGO_X)
            self.assertAlmostEqual(position(logo, 'V'), LOGO_Y)
            self.assertAlmostEqual(int(logo.find(qn('wp:extent')).get('cx'))/914400, LOGO_WIDTH)
            self.assertEqual(logo.find(qn('wp:positionH')).get('relativeFrom'), 'page')
            self.assertIn(DEFAULT_LOGO_PATH.read_bytes(), media_bytes(data))
            self.assertEqual(len(doc.tables), 0)
            self.assertEqual(len(doc.inline_shapes), 0)
            anchors = doc.element.xpath('//wp:anchor')
            if photo is None or photo == b'invalid':
                self.assertEqual(anchors, [])
                continue
            self.assertEqual(len(anchors), 2)
            image, caption = anchors
            self.assertEqual((FRAME_X, FRAME_Y, FRAME_WIDTH, FRAME_HEIGHT), (5.38, 2.29, 2.48, 3.28))
            for axis, expected in (('H', FRAME_X), ('V', FRAME_Y)):
                self.assertAlmostEqual(position(image, axis), expected)
                self.assertEqual(image.find(qn('wp:position'+axis)).get('relativeFrom'), 'page')
            extent = image.find(qn('wp:extent'))
            self.assertAlmostEqual(int(extent.get('cx'))/914400, FRAME_WIDTH)
            self.assertAlmostEqual(int(extent.get('cy'))/914400, FRAME_HEIGHT)
            self.assertEqual(image.find(qn('wp:wrapSquare')).get('wrapText'), 'left')
            outline = image.find('.//'+qn('a:ln'))
            self.assertEqual(int(outline.get('w'))/12700, 1)
            self.assertEqual(outline.find('.//'+qn('a:srgbClr')).get('val'), '141638')
            self.assertAlmostEqual(position(caption, 'V'), FRAME_Y + FRAME_HEIGHT - .18)
            paragraph = Paragraph(caption.find('.//'+qn('w:p')), doc)
            self.assertEqual(paragraph.text, 'AUDIT REFERENCE IMAGE')
            self.assertEqual(paragraph.alignment, 1)
            self.assertEqual(paragraph.runs[0].font.size.pt, 5)
            self.assertEqual(paragraph.runs[0].font.name, BOLD_FONT)
            self.assertTrue(paragraph.runs[0].italic)
            self.assertIn(photo, media_bytes(data))

    def test_constraints_release_below_panel_and_rhythm_is_count_independent(self):
        result = completed_result(4)
        for item in result.opportunities:
            item.evidence = 'The large Rollback sign dominates the pallet, while the exposed cardboard tier edges and black base carry no Tide graphics.'
            item.recommendation = 'Give the existing pallet more Tide presence with branded tier-edge strips and a base wrap, keeping the Rollback price clearly visible.'
        doc = Document(BytesIO(render_docx(build_report(result, audit_photo=photo_bytes()))))
        opportunity_copy = doc.paragraphs[5:]
        indents = [p.paragraph_format.right_indent.inches for p in opportunity_copy]
        self.assertAlmostEqual(indents[0], 2.64, places=3)
        self.assertIn(0, indents)
        first_full = indents.index(0)
        self.assertTrue(all(indent == 0 for indent in indents[first_full:]))
        top_rule = doc.paragraphs[4]._p.pPr.find(qn('w:pBdr')).find(qn('w:bottom'))
        self.assertEqual(top_rule.get(qn('w:color')), '141638')
        headings = [p for p in opportunity_copy if p.text[:2].isdigit()]
        for heading in headings[1:]:
            divider = heading._p.getprevious()
            border = divider.find(qn('w:pPr')).find(qn('w:pBdr')).find(qn('w:top'))
            self.assertEqual(border.get(qn('w:sz')), '4')
            self.assertEqual(border.get(qn('w:color')), '141638')
            self.assertEqual(heading.paragraph_format.space_before.pt, 11)
        formats = []
        for count in (1,2,3,4):
            document = Document(BytesIO(render_docx(build_report(completed_result(count)))))
            formats.append([p._p.pPr.xml for p in document.paragraphs[5:10]])
        self.assertTrue(all(fmt == formats[0] for fmt in formats))

    def test_bottom_signature_date_and_centered_footer_are_stable(self):
        zones = []
        for count in (1,2,3,4,8):
            doc = Document(BytesIO(render_docx(build_report(completed_result(count)))))
            self.assertFalse(any(p.text.startswith(('Signature:', 'Date:')) for p in doc.paragraphs))
            section = doc.sections[0]
            self.assertEqual(section.bottom_margin.inches, 1)
            table = section.footer.tables[0]
            self.assertEqual(len(table.rows), 1)
            self.assertFalse(table.autofit)
            cells = table.rows[0].cells
            self.assertEqual([cell.text for cell in cells], ['Signature:', '', '', 'Date:', ''])
            for index in (0,3):
                run = cells[index].paragraphs[0].runs[0]
                self.assertEqual(run.font.name, BLACK_FONT)
                self.assertEqual(run.font.size.pt, 10)
                self.assertEqual(cells[index].paragraphs[0].paragraph_format.line_spacing.pt, 14)
            for index in (1,4):
                self.assertAlmostEqual(cells[index].width.inches, 2.1, places=3)
                border = cells[index]._tc.tcPr.find(qn('w:tcBorders')).find(qn('w:bottom'))
                self.assertEqual(border.get(qn('w:val')), 'single')
                self.assertEqual(border.get(qn('w:sz')), '4')
                self.assertEqual(border.get(qn('w:color')), '141638')
            self.assertNotIn('_', ''.join(cell.text for cell in cells))
            self.assertNotIn('Approved By', ''.join(cell.text for cell in cells))
            spacer, separator, footer = section.footer.paragraphs
            self.assertEqual(spacer.paragraph_format.line_spacing.pt, 1)
            self.assertEqual(separator.paragraph_format.space_before.pt, 8)
            self.assertEqual(separator.paragraph_format.space_after.pt, 10)
            self.assertEqual(separator.paragraph_format.line_spacing.pt, 1)
            self.assertEqual(separator._p.pPr.find(qn('w:pBdr'))[0].get(qn('w:space')), '0')
            self.assertEqual(footer.alignment, 1)
            self.assertEqual(footer.runs[0].font.name, BOOK_FONT)
            self.assertEqual(footer.runs[0].font.size.pt, 8)
            self.assertAlmostEqual(section.footer_distance.inches, .18, places=3)
            zones.append(str(section.footer._element.xml))
        self.assertEqual(len(set(zones)), 1)

    def test_opportunity_rhythm_and_divider_gaps_are_fixed_for_all_counts(self):
        for count in (1,2,3,4):
            report = build_report(completed_result(count), audit_photo=photo_bytes())
            doc = Document(BytesIO(render_docx(report)))
            rhythms = []
            for item in report.opportunities:
                start = next(i for i,p in enumerate(doc.paragraphs) if p.text == f'{item.number:02d} {item.title}')
                paragraphs = doc.paragraphs[start:start+5]
                self.assertEqual([p.text for p in paragraphs[1::2]], ['WHY','RECOMMENDATION'])
                before = [p.paragraph_format.space_before.pt for p in paragraphs]
                after = [p.paragraph_format.space_after.pt for p in paragraphs]
                self.assertEqual(before, [11,0,0,0,0])
                self.assertEqual(after, [9,1,10,1,13])
                rhythms.append((before,after))
                if item.number > 1:
                    divider = doc.paragraphs[start-1]
                    self.assertEqual(divider.text, '')
                    self.assertEqual(divider.paragraph_format.line_spacing.pt, 1)
                    self.assertEqual(divider.paragraph_format.space_before.pt, 0)
                    self.assertEqual(divider.paragraph_format.space_after.pt, 0)
                    border = divider._p.pPr.find(qn('w:pBdr'))[0]
                    self.assertEqual(border.get(qn('w:space')), '0')
                    self.assertEqual(border.get(qn('w:sz')), '4')
                    self.assertEqual(doc.paragraphs[start-2].paragraph_format.space_after.pt, 13)
            self.assertTrue(all(rhythm == rhythms[0] for rhythm in rhythms))

    def test_missing_book_is_an_explicit_dependency_error(self):
        with patch('shelf_audit.reporting.docx_renderer.discover_font_assets', return_value={}):
            with self.assertRaisesRegex(FileNotFoundError, 'Gotham-Book.ttf'):
                render_docx(build_report(completed_result(1)))

    def test_long_reports_flow_without_type_changes_or_truncation(self):
        result = completed_result(8)
        for index, item in enumerate(result.opportunities):
            item.evidence = f'Long test section {index}. ' * 100
        report = build_report(result, audit_photo=photo_bytes())
        self.assertEqual(select_layout(report), LAYOUT_PRESETS[0])
        data = render_docx(report)
        for item in report.opportunities:
            self.assertIn(item.evidence, docx_text(data))
        doc = Document(BytesIO(data))
        self.assertEqual(doc.styles['Normal'].font.size.pt, 10)
        self.assertIn('Signature:', doc.sections[0].footer.tables[0].cell(0,0).text)


    def test_image_orientation_does_not_change_opportunity_spacing_or_estimates(self):
        for count in (2, 3, 4):
            reports = [build_report(completed_result(count), audit_photo=photo_bytes(size)) for size in ((300, 900), (900, 300))]
            self.assertEqual(estimate_report_height(reports[0], select_layout(reports[0])), estimate_report_height(reports[1], select_layout(reports[1])))
            documents = [Document(BytesIO(render_docx(report))) for report in reports]
            for first, second in zip(documents[0].paragraphs[5:], documents[1].paragraphs[5:]):
                self.assertEqual(first.text, second.text)
                self.assertEqual(first._p.pPr.xml, second._p.pPr.xml)

    def test_font_helper_sets_real_word_bold_for_gotham_bold(self):
        for family, expected in ((MEDIUM_FONT, False), (BOLD_FONT, True), (BLACK_FONT, True)):
            run = Document().add_paragraph().add_run("Sample")
            _set_font(run, family)
            self.assertEqual(run.bold, expected)
            self.assertEqual(run.font.name, family)
            for script in ("ascii", "hAnsi", "eastAsia", "cs"):
                self.assertEqual(run._r.rPr.rFonts.get(qn(f"w:{script}")), family)

    def test_private_gotham_assets_and_embedded_font_package(self):
        asset_dir = Path(__file__).resolve().parents[1] / "shelf_audit" / "assets" / "fonts"
        expected_names = {
            "Gotham Black": "Gotham-Black.ttf",
            "Gotham Bold": "Gotham-Bold.ttf",
            "Gotham Medium": "Gotham-Medium.ttf",
            "Gotham Book": "Gotham-Book.ttf",
        }
        for family_name, filename in expected_names.items():
            path = asset_dir / filename
            self.assertTrue(path.exists(), f"Missing font asset: {path}")
            self.assertTrue(path.stat().st_size > 0)
            font = TTFont(str(path))
            names = {record.toUnicode() for record in font["name"].names if record.nameID in (1, 16, 17)}
            self.assertIn(family_name, names)
        report = build_report(completed_result(1))
        package = ZipFile(BytesIO(render_docx(report)))
        names = set(package.namelist())
        self.assertIn("word/fontTable.xml", names)
        self.assertIn("word/_rels/fontTable.xml.rels", names)
        self.assertIn("word/fonts/Gotham-Medium.odttf", names)
        self.assertIn("word/fonts/Gotham-Bold.odttf", names)
        self.assertIn("word/fonts/Gotham-Black.odttf", names)
        document_rels = package.read("word/_rels/document.xml.rels").decode("utf-8")
        self.assertIn("fontTable.xml", document_rels)
        font_table = package.read("word/fontTable.xml").decode("utf-8")
        for family_name in expected_names:
            self.assertIn(family_name, font_table)
        self.assertIn("Gotham-Medium.odttf", package.read("word/_rels/fontTable.xml.rels").decode("utf-8"))
        self.assertIn("Gotham-Bold.odttf", package.read("word/_rels/fontTable.xml.rels").decode("utf-8"))
        self.assertIn("Gotham-Black.odttf", package.read("word/_rels/fontTable.xml.rels").decode("utf-8"))


    def test_obfuscated_font_keys_permissions_and_package_integrity(self):
        from shelf_audit.reporting.font_embedding import (
            discover_font_assets, inspect_font, embed_report_fonts, NS_W, NS_R, NS_PKG_REL, NS_CT,
        )
        data = render_docx(build_report(completed_result(3), audit_photo=photo_bytes()))
        with ZipFile(BytesIO(data)) as package:
            self.assertIsNone(package.testzip())
            self.assertEqual(len(package.namelist()), len(set(package.namelist())))
            self.assertFalse(any(name.endswith('.ttf') for name in package.namelist()))
            table = ET.fromstring(package.read('word/fontTable.xml'))
            # Compatibility prefixes in attribute values must retain declarations.
            from lxml import etree
            font_root = etree.fromstring(package.read('word/fontTable.xml'))
            ignorable = font_root.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable', '')
            self.assertTrue(all(prefix in font_root.nsmap for prefix in ignorable.split()))
            rels = ET.fromstring(package.read('word/_rels/fontTable.xml.rels'))
            targets = {rel.get('Id'): 'word/' + rel.get('Target') for rel in rels}
            self.assertEqual(len(targets), 4)
            self.assertEqual(len(set(targets.values())), 4)
            for rel in rels:
                self.assertEqual(rel.get('Type'), NS_R + '/font')
            types = ET.fromstring(package.read('[Content_Types].xml'))
            overrides = {item.get('PartName'): item.get('ContentType') for item in types.findall('{'+NS_CT+'}Override')}
            for family, path in discover_font_assets().items():
                info = inspect_font(path)
                self.assertEqual(info['family_name'], family)
                self.assertEqual(info['full_name'], family)
                self.assertEqual(info['fs_type'], 0x0104)
                self.assertTrue(info['embed_allowed'])
                entry = next(item for item in table if item.get('{'+NS_W+'}name') == info['family_name'])
                role = 'embedBold' if family in ('Gotham Bold', 'Gotham Black') else 'embedRegular'
                reference = entry.find('{'+NS_W+'}'+role)
                key = uuid.UUID(reference.get('{'+NS_W+'}fontKey')).bytes[::-1]
                part = targets[reference.get('{'+NS_R+'}id')]
                self.assertEqual(overrides['/'+part], 'application/vnd.openxmlformats-officedocument.obfuscatedFont')
                embedded = package.read(part)
                original = path.read_bytes()
                self.assertNotEqual(embedded[:32], original[:32])
                recovered = bytes(value ^ key[index % 16] for index, value in enumerate(embedded[:32])) + embedded[32:]
                self.assertEqual(recovered, original)
                self.assertEqual(reference.get('{'+NS_W+'}subsetted'), 'false')
            repeated = embed_report_fonts(data)
            with ZipFile(BytesIO(repeated)) as again:
                self.assertEqual(set(package.namelist()), set(again.namelist()))
                for name in package.namelist():
                    self.assertEqual(package.read(name), again.read(name), name)
            with patch('shelf_audit.reporting.docx_renderer.embed_report_fonts', side_effect=lambda value: value):
                plain = render_docx(build_report(completed_result(3), audit_photo=photo_bytes()))
            with ZipFile(BytesIO(plain)) as before:
                for name in before.namelist():
                    if name not in ('[Content_Types].xml', 'word/_rels/document.xml.rels', 'word/fontTable.xml'):
                        self.assertEqual(before.read(name), package.read(name), name)

    def test_embedding_flags_reject_restricted_and_bitmap_only_fonts(self):
        from shelf_audit.reporting.font_embedding import discover_font_assets, inspect_font, embed_report_fonts
        source = discover_font_assets()['Gotham Medium']
        for flags, allowed in ((0, True), (2, False), (4, True), (8, True), (0x104, True), (0x204, False)):
            with self.subTest(flags=flags), tempfile.TemporaryDirectory() as directory:
                target = Path(directory) / 'unrelated-filename.ttf'
                with TTFont(str(source)) as font:
                    font['OS/2'].fsType = flags
                    font.save(target)
                info = inspect_font(target)
                self.assertEqual(info['family_name'], 'Gotham Medium')
                self.assertEqual(info['fs_type'], flags)
                self.assertEqual(info['embed_allowed'], allowed)
                if not allowed:
                    plain = BytesIO()
                    Document().save(plain)
                    with self.assertLogs('shelf_audit.reporting.font_embedding', level='WARNING') as logs:
                        result = embed_report_fonts(plain.getvalue(), font_paths={'Gotham Medium': target})
                    self.assertEqual(result, plain.getvalue())
                    self.assertIn(f'fsType={flags}', '\n'.join(logs.output))

    def test_obfuscation_uses_textual_guid_reverse_order(self):
        from shelf_audit.reporting.font_embedding import _obfuscate_font
        # Fixed independent vector catches UUID little-endian/mixed-endian mistakes.
        key = '{00112233-4455-6677-8899-AABBCCDDEEFF}'
        expected = bytes.fromhex('ffeeddccbbaa99887766554433221100') * 2
        self.assertEqual(_obfuscate_font(bytes(32) + b'unchanged', key), expected + b'unchanged')

    def test_missing_invalid_photo_omits_frame_without_empty_columns(self):
        for photo in (None, b"invalid"):
            doc = Document(BytesIO(render_docx(build_report(completed_result(2), audit_photo=photo))))
            self.assertEqual(len(doc.tables), 0)
            self.assertIn("Visible evidence 0.", "\n".join(p.text for p in doc.paragraphs))

    def test_footer_year_is_calculated_at_render_time_without_mutating_report(self):
        report = build_report(completed_result(1))
        before = asdict(report)
        current_year = datetime.now(ZoneInfo("America/Chicago")).year
        self.assertIn(f"Generated by Kendal King © {current_year}", docx_text(render_docx(report)))
        for year in (2031, 2040):
            with self.subTest(year=year), patch("shelf_audit.reporting.docx_renderer.datetime") as clock:
                clock.now.return_value = datetime(year, 1, 1, tzinfo=ZoneInfo("America/Chicago"))
                self.assertIn(f"Generated by Kendal King © {year}", docx_text(render_docx(report)))
                clock.now.assert_called_once_with(ZoneInfo("America/Chicago"))
        self.assertEqual(asdict(report), before)

    def test_counts_order_verbatim_content_and_immutable_inputs(self):
        for count in (0, 1, 2, 3, 4, 8):
            with self.subTest(count=count):
                result = completed_result(count)
                report = build_report(result, audit_photo=photo_bytes())
                before_result, before_report = asdict(result), asdict(report)
                data = render_docx(report)
                text = docx_text(data)
                positions = []
                for item in report.opportunities:
                    positions.append(text.index(item.title))
                    for value in (item.title, item.evidence, item.recommendation):
                        self.assertIn(value, text)
                        self.assertEqual(text.count(value), 1)
                self.assertEqual(positions, sorted(positions))
                for hidden in (
                    "criterion_id", "confidence", "relevance", "impact", "actionability",
                    "priority_score", "provider", "prompt_version", "internal-",
                ):
                    self.assertNotIn(hidden, text)
                self.assertEqual(asdict(result), before_result)
                self.assertEqual(asdict(report), before_report)
                self.assertIn(report.footer_text, text)
                self.assertIn(report.title, text)
                self.assertIn(report.subtitle.upper(), text)
                self.assertNotIn("Approved By", text)
                self.assertIn("Signature:", text)
                self.assertIn("Date:", text)
                if count == 0:
                    self.assertIn(NO_OPPORTUNITIES_TEXT, text)

    def test_only_three_metadata_fields_rendered_even_if_others_supplied(self):
        metadata = ReportMetadata(
            client="Test Client", retailer="Test Retailer", location="Test Location",
            brand="Hidden Brand", audit_date=date(2026, 10, 12),
            account_lead="Hidden Lead", audit_reference="Hidden Audit",
            photo_reference="Hidden Photo", signature="Hidden Signature",
            signature_date=date(2026, 10, 13),
        )
        text = docx_text(render_docx(build_report(completed_result(1), metadata=metadata)))
        for value in (metadata.client, metadata.retailer, metadata.location):
            self.assertIn(value, text)
        for value in ("Hidden", "2026-10-12", "2026-10-13", "Account Lead", "Audit Date"):
            self.assertNotIn(value, text)
        blank = docx_text(render_docx(build_report(completed_result(0))))
        for label in ("Client:", "Retailer:", "Location:"):
            self.assertIn(label, blank)
        self.assertNotIn("None", blank)

    def test_literal_markup_unicode_and_whitespace_are_preserved(self):
        result = completed_result(1)
        result.opportunities[0].title = "  <b>Test title</b>  "
        result.opportunities[0].evidence = "Café — shopper’s view\nSecond line <script>literal</script>"
        result.opportunities[0].recommendation = "Keep 100% of text ©.\tNext step."
        report = build_report(result)
        text = docx_text(render_docx(report))
        for item in report.opportunities:
            for value in (item.title, item.evidence, item.recommendation):
                self.assertIn(value, text)

    def test_original_photo_bytes_and_bounded_aspect_ratio(self):
        for size, image_format in (((900, 300), "PNG"), ((300, 900), "PNG"), ((400, 300), "JPEG")):
            with self.subTest(size=size, image_format=image_format):
                photo = photo_bytes(size, image_format=image_format)
                report = build_report(completed_result(2), audit_photo=photo)
                self.assertIs(report.audit_photo, photo)
                data = render_docx(report)
                self.assertIn(photo, media_bytes(data))
                self.assertIn(DEFAULT_LOGO_PATH.read_bytes(), media_bytes(data))
                doc = Document(BytesIO(data))
                extent = doc.element.xpath("//wp:anchor/wp:extent")[0]
                fill = doc.element.xpath("//wp:anchor//a:fillRect")[0]
                width = int(extent.get("cx")) / 914400 * (1 - (int(fill.get("l")) + int(fill.get("r"))) / 100000)
                height = int(extent.get("cy")) / 914400 * (1 - (int(fill.get("t")) + int(fill.get("b"))) / 100000)
                self.assertAlmostEqual(width / height, size[0] / size[1], delta=0.0002)
                self.assertLessEqual(width, PHOTO_MAX_WIDTH + 0.0001)
                self.assertLessEqual(height, PHOTO_MAX_HEIGHT + 0.0001)
                self.assertEqual(fill.get("l"), fill.get("r"))
                self.assertEqual(fill.get("t"), fill.get("b"))

    def test_missing_invalid_photo_and_missing_logo_degrade_gracefully(self):
        with tempfile.TemporaryDirectory() as directory:
            for photo in (None, b"", b"not an image"):
                with self.subTest(photo=photo):
                    data = render_docx(build_report(completed_result(1), audit_photo=photo), logo_path=Path(directory) / "missing.png")
                    self.assertIn("Visible evidence 0.", docx_text(data))
                    self.assertEqual(media_bytes(data), [])
                    self.assertFalse(Document(BytesIO(data)).element.xpath("//wp:anchor"))

    def test_short_fixtures_fit_estimated_one_page_without_forced_breaks(self):
        for count, size in ((count, size) for count in (1, 2, 3, 4) for size in ((300, 900), (900, 300))):
            with self.subTest(count=count, size=size):
                report = build_report(completed_result(count), audit_photo=photo_bytes(size))
                layout = select_layout(report)
                available = (11 - HEADER_TOP_MARGIN - layout.margin) * 72 - 24
                self.assertLessEqual(estimate_report_height(report, layout), available)
                doc = Document(BytesIO(render_docx(report)))
                self.assertFalse(doc.element.xpath('//w:br[@w:type="page"]'))
                self.assertEqual(len(doc.sections), 1)
                self.assertGreaterEqual(doc.styles["Normal"].font.size.pt, MIN_BODY_FONT_SIZE)

    def test_docx_ui_does_not_import_audit_provider_pdf_or_preview(self):
        script = '''
import sys
from shelf_audit.reporting.models import ShelfOpportunityReport
from shelf_audit.reporting.docx_renderer import render_docx
from shelf_audit.reporting.ui import render_report_section
render_docx(ShelfOpportunityReport())
assert not any(name.startswith('shelf_audit.providers') for name in sys.modules)
for name in ('shelf_audit.audit', 'shelf_audit.reporting.pdf_renderer', 'shelf_audit.reporting.preview_renderer'):
    assert name not in sys.modules, name
'''
        result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)


class ShelfAuditDocxUITests(unittest.TestCase):
    def test_three_fields_one_download_no_preview_and_no_provider_calls(self):
        import streamlit as st
        result = completed_result(2)
        app = AppTest.from_string('''
from shelf_audit.page import render
render()
''')
        app.session_state["shelf_audit_completed_result"] = result
        with (
            patch("shelf_audit.page.run_audit", side_effect=AssertionError("audit rerun")),
            patch("shelf_audit.page.OpenAIAuditProvider", side_effect=AssertionError("provider constructed")),
            patch("shelf_audit.reporting.ui.st.download_button", wraps=st.download_button) as download,
            patch("shelf_audit.reporting.ui.st.html", side_effect=AssertionError("preview rendered")),
        ):
            app.run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual([field.label for field in app.text_input], ["Client", "Retailer", "Location"])
            self.assertEqual(len(app.date_input), 0)
            self.assertEqual(len(app.get("download_button")), 1)
            self.assertEqual(len(app.button), 0)
            self.assertEqual(download.call_args.args[0], "Download Report DOCX")
            self.assertEqual(download.call_args.kwargs["on_click"], "ignore")
            self.assertTrue(download.call_args.kwargs["file_name"].endswith(".docx"))
            app.text_input(key="shelf_report_client").set_value("Test Client").run()
            self.assertEqual(len(app.exception), 0)
            self.assertIn("Test Client", docx_text(download.call_args.kwargs["data"]))
            app.run()
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any(element.value == "Shelf Opportunities" for element in app.subheader))

    def test_completed_audit_passes_original_photo_then_controls_never_reaudit(self):
        import streamlit as st
        original = photo_bytes((800, 300), "green")
        replacement = photo_bytes((300, 800), "yellow")
        uploaded = Mock()
        uploaded.name, uploaded.type = "original.png", "image/png"
        uploaded.getvalue.return_value = original
        first, second = completed_result(2), completed_result(1)
        app = AppTest.from_string('''
from shelf_audit.page import render
render()
''')
        with (
            patch("shelf_audit.page.st.file_uploader", return_value=uploaded),
            patch("shelf_audit.page.run_audit", side_effect=[first, second]) as audit,
            patch("shelf_audit.page.OpenAIAuditProvider") as provider,
            patch("shelf_audit.reporting.ui.st.download_button", wraps=st.download_button) as download,
        ):
            app.run()
            app.radio(key="shelf_audit_image_source").set_value("Upload Existing Photo").run()
            app.button(key="shelf_audit_analyze").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(audit.call_count, 1)
            self.assertEqual(app.session_state["shelf_audit_completed_photo"], original)
            self.assertIn(original, media_bytes(download.call_args.kwargs["data"]))
            # A replacement input invalidates the prior result and download;
            # the next explicit analysis associates its own original photo.
            uploaded.getvalue.return_value = replacement
            app.text_input(key="shelf_report_location").set_value("Test Store").run()
            self.assertEqual(len(app.exception), 0)
            self.assertNotIn("shelf_audit_completed_result", app.session_state)
            self.assertNotIn("shelf_audit_completed_photo", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)
            self.assertEqual(audit.call_count, 1)
            self.assertEqual(provider.call_count, 1)
            provider.return_value.analyze.assert_not_called()
            app.button(key="shelf_audit_analyze").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(audit.call_count, 2)
            self.assertIn(replacement, media_bytes(download.call_args.kwargs["data"]))
            self.assertEqual(app.session_state["shelf_audit_completed_result"], second)

    def test_failed_replacement_audit_clears_previous_photo_and_result(self):
        uploaded = Mock()
        uploaded.name, uploaded.type = "shelf.png", "image/png"
        uploaded.getvalue.return_value = photo_bytes()
        app = AppTest.from_string('''
from shelf_audit.page import render
render()
''')
        app.session_state["shelf_audit_completed_result"] = completed_result(1)
        app.session_state["shelf_audit_completed_photo"] = photo_bytes()
        with (
            patch("shelf_audit.page.st.camera_input", return_value=uploaded),
            patch("shelf_audit.page.run_audit", side_effect=RuntimeError("test failure")),
            patch("shelf_audit.page.OpenAIAuditProvider"),
        ):
            app.run()
            app.button(key="shelf_audit_analyze").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertNotIn("shelf_audit_completed_photo", app.session_state)
            self.assertNotIn("shelf_audit_completed_result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)

    def test_blank_fields_zero_findings_and_missing_photo_can_download(self):
        import streamlit as st
        app = AppTest.from_string('''
from shelf_audit.reporting.ui import render_report_section
from shelf_audit.models import AuditResult, AuditMode, AuditDepth
render_report_section(AuditResult([], AuditMode.PRODUCT, AuditDepth.DEEP, 0))
''')
        with patch("shelf_audit.reporting.ui.st.download_button", wraps=st.download_button) as download:
            app.run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual([field.label for field in app.text_input], ["Client", "Retailer", "Location"])
            self.assertIn(NO_OPPORTUNITIES_TEXT, docx_text(download.call_args.kwargs["data"]))


if __name__ == "__main__":
    unittest.main()
